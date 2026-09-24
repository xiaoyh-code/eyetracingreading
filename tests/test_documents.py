import io

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from reader import demo
from reader.documents import MAX_CHARACTERS, DocumentError, extract_upload, make_document


def plain_text(document):
    return "\n\n".join("".join(sentence["text"] for sentence in paragraph["sentences"]) for paragraph in document["paragraphs"])


def pdf_bytes(text: str | None = None) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    if text:
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def test_tokens_preserve_whitespace_punctuation_and_have_unique_ids():
    text = "Dr. Green can't wait—really!  She said, ‘It's fine.’\nNext line.\n\nA second paragraph: 3.14, naïve, 中文."
    document = make_document(text, "Example")
    assert plain_text(document) == text
    identifiers = [document["id"]]
    words = []
    for paragraph in document["paragraphs"]:
        identifiers.append(paragraph["id"])
        for sentence in paragraph["sentences"]:
            identifiers.append(sentence["id"])
            assert "".join(token["text"] for token in sentence["tokens"]) == sentence["text"]
            identifiers.extend(token["id"] for token in sentence["tokens"])
            words.extend(token["text"] for token in sentence["tokens"] if token["word"])
    assert len(set(identifiers)) == len(identifiers)
    assert "can't" in words
    assert document["word_count"] == len(words)
    assert len(document["paragraphs"]) == 2


def test_markdown_keeps_readable_text_without_markup_or_urls():
    document = extract_upload(b"# A heading\n\nA **bold** [link](https://example.com) and `code`.\n\n<script>evil()</script>\n\n![A fox](https://example.com/fox.png)", "readme.md")
    extracted = plain_text(document)
    assert "A heading" in extracted
    assert "A bold link and code." in extracted
    assert "A fox" in extracted
    assert "https://" not in extracted
    assert "<script>" not in extracted
    assert "evil()" not in extracted


def test_pdf_text_extraction_and_empty_scan_error():
    document = extract_upload(pdf_bytes("A river moves slowly."), "article.pdf")
    assert "A river moves slowly." in plain_text(document)
    assert document["source"] == "pdf"
    assert document["warnings"]
    with pytest.raises(DocumentError, match="OCR"):
        extract_upload(pdf_bytes(), "scanned.pdf")


def test_docx_body_and_table_order():
    document = Document()
    document.add_paragraph("Before the table.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Left cell"
    table.cell(0, 1).text = "Right cell"
    document.add_paragraph("After the table.")
    buffer = io.BytesIO()
    document.save(buffer)
    extracted = plain_text(extract_upload(buffer.getvalue(), "notes.docx"))
    assert extracted.index("Before") < extracted.index("Left cell") < extracted.index("After")
    assert "Right cell" in extracted


@pytest.mark.parametrize(("data", "filename", "match"), [
    (b"", "empty.txt", "空白"),
    (b"hello", "old.doc", "DOCX"),
    (b"not a PDF", "broken.pdf", "未能讀取"),
    (b"\xff\x00\x12", "binary.txt", "編碼"),
])
def test_import_errors_are_useful(data, filename, match):
    with pytest.raises(DocumentError, match=match):
        extract_upload(data, filename)


def test_document_size_is_bounded():
    with pytest.raises(DocumentError) as error:
        make_document("a" * (MAX_CHARACTERS + 1))
    assert error.value.status_code == 413


def test_demo_has_handwritten_translation_for_every_sentence():
    document = make_document(demo.TEXT)
    sentences = [sentence for paragraph in document["paragraphs"] for sentence in paragraph["sentences"]]
    assert 350 < document["word_count"] < 500
    assert all(sentence["text"].strip() in demo.TRANSLATIONS for sentence in sentences)
    assert {"serendipity", "resilience", "ephemeral"} <= demo.DICTIONARY.keys()
